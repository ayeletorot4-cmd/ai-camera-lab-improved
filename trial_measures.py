"""Reaction time, movement time etc. per trial - the same method as the previous (MATLAB) system - and the
Excel output.

Port of loaddata.m + calculateMeasures.m + makeCSVsubject.m / makeCSVall.m:
  - each trial's data runs from the end of the previous trial (the home phase) to the end of this trial
  - the marker (wrist for the upper limb, ankle for the lower limb - 3D with two cameras, otherwise its
    image position) is resampled to 100 Hz, low-pass filtered (Butterworth, 2nd order, 10 Hz, zero-phase),
    differentiated, and the tangential velocity is the speed
  - first peak = the first speed peak (> 30% of the maximum after the light came on) after the light came on
  - movement onset = the last sample between the light coming on and the first peak with speed < 5% of the
    maximum (calculateMeasures.m). If the speed never drops below 5% there (the resting hand drifting slightly,
    or the end of the return home) calculateMeasures.m fell back to the light-on sample, giving RT ~ 0; here the
    same 5% rule is then measured from the resting level instead of from 0: the last sample with speed <= the
    lowest speed between the light and the peak + 5% of the maximum -> reaction time (responsetime) = onset -
    light on
  - negative reaction times: if the reach toward the target already started before the light (the same < 5%
    search, but looking back before the light, at most ANTICIPATION_MAX_S), that earlier sample is the onset
    and the reaction time is negative (the subject knew the next light of the sequence). The data of each
    trial include the home phase before the light for this (as loaddata.m: "so we can include negative onset
    times").
  - movement end = the first sample after the first peak with speed < 5% of the maximum (the last sample if
    there is none) -> movement time = end - onset
  - also: straight-line deviation, number of speed peaks (> 10% of the maximum), acceleration time (AT),
    deceleration time (DT), AT/DT, end-point distance to the target
  - summaries: mean, median, SD and IQR by part, by block and by sequence, without familiarization, only
    the trials with touse = 1 (MATLAB's std and iqr: sample SD, 'hazen' percentiles)
Flags (new):
  - Anticipation: the movement toward the target started before the light came on (negative reaction time)
  - Not from home (shortcut): the movement didn't start at the home position (where the hand/foot rests when
    the lights come on) but e.g. at the previous switch - with where it started
  - Corrected: the first movement went to another switch, then the right one was hit in time
"""
import csv
import os
import warnings

import numpy as np
from scipy.signal import butter, filtfilt, find_peaks

RESAMPLE_HZ = 100
FILTER_B, FILTER_A = butter(2, 10 / (RESAMPLE_HZ / 2))
MARKER = {'upper': {'right': 16, 'left': 15}, 'lower': {'right': 28, 'left': 27}}  # wrist / ankle
MIN_VISIBILITY = 0.3
ANTICIPATION_MIN_S = 0.03        # movement onset at least this much before the light -> anticipation
ANTICIPATION_MAX_S = 0.5         # ... and at most this much (earlier = still the return to home, not a reach)
TOWARD_COS = 0.7                 # 'toward the target' = moving within ~45 degrees of the direction to it
HOME_ZONE = 0.5                  # started more than halfway from home toward the target = not from home (the
                                 # resting position itself varies by up to ~40% of the reach)
SWITCH_ZONE = 0.25               # ... and started 'at switch k' if within 25% of the reach from that switch
CORRECTION_MIN_S = 0.15          # first movement stopped at least this long before the hit ...
EXCLUDED_FROM_SUMMARY = ['Familiarization']
SUMMARY_PARTS = ['pre_test', 'pre_test_contralateral', 'practice', 'post_test', 'post_test_contralateral',
                 're_test', 're_test_contralateral']
MEASURES = ['responsetime', 'movementtime', 'straighlinedeviation', 'numpeaks', 'AT', 'DT', 'AT_over_DT',
            'endpointdistance']


SPIKE_FRACTION = 0.03            # a point > 3% of the image size away from its neighbours' median is a glitch


def marker_track(pose, limb, side, image_size=None):
    """(frames, 2) marker position (NaN where MediaPipe isn't sure, or where it jumps for a single frame or two
    - a glitch, not a movement: compared with the median of the 5 frames around it)."""
    j = MARKER[limb][side]
    xy = pose[:, j, :2].astype(float).copy()
    xy[~(pose[:, j, 2] >= MIN_VISIBILITY)] = np.nan
    if image_size is not None and len(xy) >= 5:
        pad = np.pad(xy, ((2, 2), (0, 0)), mode='edge')
        with warnings.catch_warnings():   # all-NaN stretches (marker not seen) are expected
            warnings.simplefilter('ignore', RuntimeWarning)
            med = np.nanmedian(np.lib.stride_tricks.sliding_window_view(pad, 5, axis=0), axis=2)
        jump = np.linalg.norm(xy - med, axis=1) > SPIKE_FRACTION * max(image_size)
        xy[jump] = np.nan
    return xy


def _point_line_distance(pts, v1, v2):
    """point_to_line_distance.m: distance of each point to the line through v1 and v2 (2D or 3D)."""
    pad = lambda a: np.hstack([a, np.zeros((a.shape[0], 3 - a.shape[1]))]) if a.shape[1] == 2 else a
    p, a, b = pad(np.atleast_2d(pts)), pad(np.atleast_2d(v1)), pad(np.atleast_2d(v2))
    d = a - b
    n = np.linalg.norm(d)
    return np.linalg.norm(np.cross(d, p - b), axis=1) / n if n > 0 else np.linalg.norm(p - b, axis=1)


def trial_kinematics(t_s, pos, stim_s):
    """loaddata.m for one trial: 100 Hz time, position, filtered speed. None if too little data."""
    ok = ~np.isnan(pos[:, 0])
    if ok.sum() < 5 or ok.mean() < 0.5:
        return None
    time = np.arange(t_s[0], t_s[-1] + 1e-9, 1 / RESAMPLE_HZ)
    xyz = np.column_stack([np.interp(time, t_s[ok], pos[ok, c]) for c in range(pos.shape[1])])
    padlen = 3 * (max(len(FILTER_A), len(FILTER_B)) - 1)
    if len(time) <= padlen:
        return None
    filt = filtfilt(FILTER_B, FILTER_A, xyz, axis=0, padlen=padlen)
    vel = np.vstack([np.zeros((1, xyz.shape[1])), np.diff(filt, axis=0) * RESAMPLE_HZ])
    return dict(time=time, xyz=xyz, tangvel=np.linalg.norm(vel, axis=1), vel=vel, stim=stim_s)


def calculate_measures(k, target_pos, home_pos=None):
    """calculateMeasures.m for one trial (k = trial_kinematics output), with negative reaction times when the
    reach started before the light. Returns measures + plot samples."""
    nan = dict.fromkeys(MEASURES, np.nan)
    empty = dict(nan, onset_s=np.nan, end_s=np.nan, anticipation=False, anticipation_ms=None, start_pos=None)
    if k is None:
        return empty
    time, xyz, tv, stim = k['time'], k['xyz'], k['tangvel'], k['stim']
    s = int(np.searchsorted(time, stim))   # MATLAB: find(time >= stimulusonsettime, 1)
    if s >= len(time) - 1:
        return empty
    maxtv = tv[s:].max()
    peaks, _ = find_peaks(tv, height=0.3 * maxtv)
    after = peaks[time[peaks] >= stim]
    if maxtv <= 0 or len(after) == 0:
        return empty
    first = int(after[0])
    seg_tv = tv[s:first + 1]
    below = np.nonzero(seg_tv < 0.05 * maxtv)[0]
    if len(below):
        onset = s + int(below[-1])                                   # calculateMeasures.m
    else:                                                            # never < 5%: 5% above the resting level
        onset = s + int(np.nonzero(seg_tv <= seg_tv.min() + 0.05 * maxtv)[0][-1])
    # Start of the rise into the first peak (back from the peak while the speed keeps falling, until < 5%) -
    # only to tell whether the reach began before the light (a slow drift is not a rise)
    rise = first
    while rise > 0 and tv[rise - 1] < tv[rise] and tv[rise] >= 0.05 * maxtv:
        rise -= 1
    # Started before the light? It counts if the rise began at most ANTICIPATION_MAX_S before the light and the
    # hand/foot was then moving toward the target (not e.g. still returning home from the previous trial); then
    # the onset is that earlier moment and the reaction time is negative.
    anticipation = False
    lead = stim - time[rise]
    if ANTICIPATION_MIN_S < lead <= ANTICIPATION_MAX_S and target_pos is not None:
        v = k['vel'][s]
        to_target = np.asarray(target_pos, float) - xyz[s]
        nv, nt = np.linalg.norm(v), np.linalg.norm(to_target)
        if nv > 0.05 * maxtv and nt > 0 and np.dot(v, to_target) / (nv * nt) > TOWARD_COS:
            anticipation = True
            onset = rise

    below = np.nonzero(tv[first:] < 0.05 * maxtv)[0]
    end = first + int(below[0]) if len(below) else len(time) - 1
    rng = slice(onset, end + 1)
    seg = xyz[rng]
    m = dict(responsetime=time[onset] - stim, movementtime=time[end] - time[onset],
             straighlinedeviation=float(np.mean(_point_line_distance(seg, seg[0], seg[-1]))),
             numpeaks=len(find_peaks(tv[rng], height=0.1 * tv.max())[0]),
             AT=time[first] - time[onset], DT=time[end] - time[first])
    m['AT_over_DT'] = m['AT'] / m['DT'] if m['DT'] > 0 else np.nan
    m['endpointdistance'] = float(np.linalg.norm(seg[-1] - target_pos)) if target_pos is not None else np.nan
    m.update(onset_s=time[onset], end_s=time[end], peak_s=time[first], start_pos=xyz[onset],
             anticipation=anticipation, anticipation_ms=round(lead * 1000, 1) if anticipation else None)
    return m


def start_place(start, home, target_pos, switch_pos):
    """Where the movement started: ('home' | 'switch k' | 'elsewhere', distance from home in % of the
    home-target distance)."""
    if start is None or home is None or target_pos is None:
        return '', None
    reach = np.linalg.norm(np.asarray(target_pos) - home)
    if reach <= 0:
        return '', None
    d = np.linalg.norm(start - home) / reach
    if d <= HOME_ZONE:
        return 'home', round(100 * d)
    to_sw = np.linalg.norm(np.asarray(switch_pos) - start, axis=1) / reach
    k = int(np.argmin(to_sw))
    return (f'switch {k + 1}' if to_sw[k] <= SWITCH_ZONE else 'elsewhere'), round(100 * d)


def nearest_switch(p, switch_pos, switch_size):
    d = np.linalg.norm(switch_pos - p, axis=1) / switch_size
    return int(np.argmin(d)) + 1


def corrected_reach(k, m, hit_s, target, switch_pos, switch_size):
    """Which other switch the first movement went to, if the target was then hit after a correction."""
    if k is None or np.isnan(m.get('end_s', np.nan)) or hit_s is None:
        return None
    time, xyz = k['time'], k['xyz']
    first_end = nearest_switch(xyz[np.searchsorted(time, m['end_s'])], switch_pos, switch_size)
    at_hit = nearest_switch(xyz[min(len(time) - 1, np.searchsorted(time, hit_s))], switch_pos, switch_size)
    if hit_s - m['end_s'] >= CORRECTION_MIN_S and first_end != at_hit and first_end != target:
        return first_end
    return None


def _homes(trials, track):
    """Home position per block and limb/side: where the marker is when the lights come on in that block (median
    over its trials) - the resting position can shift a little between blocks / parts."""
    homes = {}
    groups = {}
    for r in trials:
        if r.get('on_frame'):
            groups.setdefault((r.get('block_index'), r['Limb'], r['Side']), []).append(r['on_frame'])
    for (bi, limb, side), frames in groups.items():
        tr = track.get((limb, side))
        if tr is None:
            continue
        pts = tr[frames]
        pts = pts[~np.isnan(pts[:, 0])]
        if len(pts) >= 5:
            homes[(bi, limb, side)] = np.median(pts, axis=0)
    return homes


def measure_trials(trials, times_ms, sources):
    """Add measures and flags to the trial rows (in place).
    sources: list of dict(name, track={(limb, side): (frames, 2 or 3) marker positions}, switch_pos (switches,
    dims), switch_size) - e.g. 3D, camera A, camera B. For each trial the first source is used if it gives a
    result, else the one that saw the marker best during the trial. Distances (straightness, end point) are
    only kept when the first source was used (other sources have other units)."""
    t_s = np.asarray(times_ms) / 1000
    homes = [_homes(trials, src['track']) for src in sources]
    for r in trials:
        r.update(dict.fromkeys(MEASURES, np.nan), anticipation=False, anticipation_ms=None, corrected_from=None,
                 touse=r['Outcome'] == 'hit', Data='', start_from='', start_dist_pct=None)
        if r['on_frame'] is None or r.get('off_frame') is None:
            r['touse'] = False
            continue
        key = (r['Limb'], r['Side'])
        a, b = r['window_start'], r['off_frame']
        cover = [np.mean(~np.isnan(src['track'][key][a:b + 1, 0])) if key in src['track'] else 0 for src in sources]
        order = [0] + sorted(range(1, len(sources)), key=lambda i: -cover[i])
        k = m = None
        for i in order:
            if cover[i] == 0:
                continue
            src = sources[i]
            k = trial_kinematics(t_s[a:b + 1], src['track'][key][a:b + 1], t_s[r['on_frame']])
            m = calculate_measures(k, src['switch_pos'][r['Target'] - 1], homes[i].get((r.get('block_index'),) + key))
            if not np.isnan(m['responsetime']):
                r['Data'] = src['name']
                break
        if m is None or not r['Data']:
            r['touse'] = False
            continue
        if r['Data'] != sources[0]['name']:
            m['straighlinedeviation'] = m['endpointdistance'] = np.nan
        r.update({x: m.get(x, np.nan) for x in MEASURES})
        r['anticipation'], r['anticipation_ms'] = m['anticipation'], m['anticipation_ms']
        r['onset_s'], r['end_s'] = m.get('onset_s'), m.get('end_s')
        r['start_from'], r['start_dist_pct'] = start_place(m.get('start_pos'), homes[i].get((r.get('block_index'),) + key),
                                                           src['switch_pos'][r['Target'] - 1], src['switch_pos'])
        if r['Outcome'] == 'hit':
            r['corrected_from'] = corrected_reach(k, m, t_s[r['off_frame']], r['Target'], src['switch_pos'],
                                                  src['switch_size'])
    return trials


# -----------------------------
# Output files
# -----------------------------
TRIAL_COLUMNS = ['Part', 'Block', 'Sequence', 'Trial', 'Target', 'Limb', 'Side', 'Light_on_ms', 'Touch_ms',
                 'Hit_ms', 'Light_off_ms', 'Trial_end_ms', 'Outcome', 'Light_source', 'Anticipation',
                 'Anticipation_ms', 'Not_from_home', 'Started_from', 'Start_distance_pct', 'Corrected',
                 'Corrected_from', 'RT_s', 'MT_s']


def _cell(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ''
    if isinstance(v, (bool, np.bool_)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return round(float(v), 4)
    if isinstance(v, np.integer):
        return int(v)
    return v


def write_trials_csv(path, trials):
    """The trials synced with the experimental design, with the flags (one row per trial)."""
    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(TRIAL_COLUMNS)
        for r in trials:
            w.writerow([_cell(r.get(c)) for c in TRIAL_COLUMNS[:14]] +
                       [_cell(r.get('anticipation')), _cell(r.get('anticipation_ms')),
                        int(r.get('start_from') not in ('', 'home')), r.get('start_from', ''),
                        _cell(r.get('start_dist_pct')),
                        int(r.get('corrected_from') is not None), _cell(r.get('corrected_from')),
                        _cell(r.get('responsetime')), _cell(r.get('movementtime'))])


def _iqr(x):
    return np.percentile(x, 75, method='hazen') - np.percentile(x, 25, method='hazen') if len(x) else np.nan


STATS = {'mean': np.mean, 'median': np.median, 'sd': lambda x: np.std(x, ddof=1) if len(x) > 1 else np.nan,
         'iqr': _iqr}


def summary_tables(rows_by_subject):
    """makeCSVall.m: 12 wide tables (mean/median/sd/iqr x by part/block/sequence), one row per subject.
    rows_by_subject: {subject: [measure rows with Part/Block/Sequence/touse + MEASURES]}."""
    tables = {f'{stat}_by{level}': [] for level in ('part', 'block', 'sequence') for stat in STATS}
    for subject, rows in rows_by_subject.items():
        use = [r for r in rows if r['touse'] and r['Part'] not in EXCLUDED_FROM_SUMMARY]
        for stat, fn in STATS.items():
            out = {lvl: {'subject': subject} for lvl in ('part', 'block', 'sequence')}
            for meas in MEASURES:
                for part in SUMMARY_PARTS:
                    pr = [r for r in use if r['Part'] == part]
                    if not pr:
                        continue
                    vals = lambda rr: np.array([r[meas] for r in rr if not np.isnan(r[meas])], float)
                    v = vals(pr)
                    out['part'][f'{stat}_{part}_{meas}'] = fn(v) if len(v) else np.nan
                    for b in sorted({r['Block'] for r in pr}):
                        br = [r for r in pr if r['Block'] == b]
                        v = vals(br)
                        out['block'][f'{stat}_{part}_block{b}_{meas}'] = fn(v) if len(v) else np.nan
                        for sq in sorted({r['Sequence'] for r in br}):
                            v = vals([r for r in br if r['Sequence'] == sq])
                            out['sequence'][f'{stat}_{part}_block{b}_sequence{sq}_{meas}'] = fn(v) if len(v) else np.nan
            for lvl in out:
                tables[f'{stat}_by{lvl}'].append(out[lvl])
    return tables


MEASURE_COLUMNS = (['Subject', 'Session', 'Part', 'Block', 'Sequence', 'Trial', 'Target', 'Outcome', 'Light_source',
                    'Anticipation',
                    'Anticipation_ms', 'Not_from_home', 'Started_from', 'Start_distance_pct', 'Corrected',
                    'Corrected_from', 'Data'] + MEASURES + ['touse'])


def measure_row(subject, session, r):
    return {'Subject': subject, 'Session': session, 'Part': r['Part'], 'Block': r['Block'],
            'Sequence': r['Sequence'], 'Trial': r['Trial'], 'Target': r['Target'], 'Outcome': r['Outcome'],
            'Light_source': r.get('Light_source', ''),
            'Anticipation': bool(r.get('anticipation')), 'Anticipation_ms': r.get('anticipation_ms'),
            'Not_from_home': r.get('start_from') not in ('', 'home', None), 'Started_from': r.get('start_from', ''),
            'Start_distance_pct': r.get('start_dist_pct'),
            'Corrected': r.get('corrected_from') is not None, 'Corrected_from': r.get('corrected_from'),
            'Data': r.get('Data', ''),
            **{m: r.get(m, np.nan) for m in MEASURES}, 'touse': bool(r.get('touse'))}


def write_analysis_xlsx(path, rows_by_subject, info=None, units=''):
    """Excel workbook: 'measures' (one row per trial) + the 12 summary sheets (+ 'info')."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'measures'
    ws.append(MEASURE_COLUMNS)
    for rows in rows_by_subject.values():
        for r in rows:
            ws.append([_cell(r.get(c)) for c in MEASURE_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for i in range(1, len(MEASURE_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    for name, table in summary_tables(rows_by_subject).items():
        sh = wb.create_sheet(name)
        cols = ['subject'] + [k for k in dict.fromkeys(k for row in table for k in row) if k != 'subject']
        sh.append(cols)
        for row in table:
            sh.append([_cell(row.get(c, np.nan)) for c in cols])
        for c in sh[1]:
            c.font = Font(bold=True)
        sh.freeze_panes = 'B2'
    sh = wb.create_sheet('info')
    lines = [('Times', 'seconds (responsetime, movementtime, AT, DT)'),
             ('Distances', units or 'see the analysis settings'),
             ('Method', 'Same as the previous MATLAB system (calculateMeasures.m): 100 Hz, 10 Hz Butterworth, '
                        '5% / 30% of peak speed thresholds'),
             ('Summaries', 'Only rows with touse = 1, familiarization excluded; sd = sample SD, iqr = MATLAB iqr')]
    for k, v in (info or {}).items():
        lines.append((k, str(v)))
    for line in lines:
        sh.append(list(line))
    sh.column_dimensions['A'].width = 18
    sh.column_dimensions['B'].width = 100
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
